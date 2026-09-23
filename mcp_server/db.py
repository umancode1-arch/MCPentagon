from __future__ import annotations

from pathlib import Path
import xml.etree.ElementTree as ET

from sqlalchemy import create_engine, delete, select
from sqlalchemy.orm import DeclarativeBase, Mapped, Session, mapped_column
from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "smartcomms.db"
TEMPLATE_ROOT = BASE_DIR / "templates"
GENERATED_ROOT = BASE_DIR / "generated"


class Base(DeclarativeBase):
    pass


class ClientTemplateMap(Base):
    __tablename__ = "client_template_map"

    id: Mapped[int] = mapped_column(primary_key=True)
    client_id: Mapped[str]
    template_name: Mapped[str]
    docx_filename: Mapped[str]


engine = create_engine(f"sqlite:///{DB_PATH}", connect_args={"check_same_thread": False})
Base.metadata.create_all(engine)


def ensure_directories() -> None:
    GENERATED_ROOT.mkdir(exist_ok=True)
    for template_name in ["account_agreement", "dapply", "dapproval"]:
        (TEMPLATE_ROOT / template_name / "docs").mkdir(parents=True, exist_ok=True)
        (TEMPLATE_ROOT / template_name / "variables").mkdir(parents=True, exist_ok=True)


def _create_template_doc(doc_path: Path, template_name: str) -> None:
    if doc_path.exists():
        return
    doc = Document()
    doc.add_heading(template_name.replace("_", " ").title(), level=1)
    doc.add_paragraph("This template is generated for demo use.")
    doc.add_paragraph("Rates and Fees Table")
    doc.add_paragraph("{{apr_rate}}")
    doc.add_paragraph("How Interest Is Calculated")
    doc.add_paragraph("{{interest_formula}}")
    doc.add_paragraph("How Fees Work")
    doc.add_paragraph("{{fee_schedule}}")
    doc.add_paragraph("Standard Provisions")
    doc.add_paragraph("{{effective_date}}")
    rows = [
        ["Fee Type", "Amount", "Description"],
        ["{{fee_type}}", "{{fee_amount}}", "{{fee_notes}}"],
    ]
    table = doc.add_table(rows=2, cols=3)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            table.cell(r_idx, c_idx).text = value
    doc.save(doc_path)


def _write_variable_xml(path: Path, template_name: str, client_id: str, version: int, values: dict[str, str]) -> None:
    root = ET.Element("variables")
    root.set("client_id", client_id)
    root.set("template_name", template_name)
    root.set("version", str(version))
    for key, value in values.items():
        elem = ET.SubElement(root, "field", {"name": key})
        elem.text = value
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


def seed_mock_data() -> None:
    ensure_directories()
    template_clients = {
        "account_agreement": ["3227", "4410", "9001"],
        "dapply": ["3227", "4410"],
        "dapproval": ["3227", "9001"],
    }
    sample_values = {
        "account_agreement": {
            "3227": {"apr_rate": "8.75%", "interest_formula": "Interest accrues on the daily ending balance at 8.75% APR.", "fee_schedule": "Monthly maintenance fee of $12 for balances below $500.", "effective_date": "2026-09-23", "fee_type": "Maintenance Fee", "fee_amount": "$12", "fee_notes": "Applies when balance is under $500.", "borrower_name": "Alicia Morgan", "account_number": "3227-1001"},
            "4410": {"apr_rate": "9.25%", "interest_formula": "Interest accrues on the average daily balance at 9.25% APR.", "fee_schedule": "Annual renewal fee of $25.", "effective_date": "2026-11-01", "fee_type": "Annual Fee", "fee_amount": "$25", "fee_notes": "Charged at renewal each calendar year.", "borrower_name": "Jordan Lee", "account_number": "4410-2044"},
            "9001": {"apr_rate": "7.90%", "interest_formula": "Interest accrues on outstanding principal using the promotional rate schedule.", "fee_schedule": "No monthly maintenance fee for preferred clients.", "effective_date": "2026-10-15", "fee_type": "Preferred Fee", "fee_amount": "$0", "fee_notes": "No charge while active.", "borrower_name": "Samira Patel", "account_number": "9001-8802"},
        },
        "dapply": {
            "3227": {"applicant_name": "Alicia Morgan", "loan_amount": "$15,000", "purpose": "Home improvement", "decision_date": "2026-09-25", "risk_score": "74", "review_note": "Strong candidate with stable income history.", "fee_type": "Application Fee", "fee_amount": "$250"},
            "4410": {"applicant_name": "Jordan Lee", "loan_amount": "$24,000", "purpose": "Vehicle purchase", "decision_date": "2026-10-02", "risk_score": "67", "review_note": "Moderate risk; requires standard disclosures.", "fee_type": "Doc Prep Fee", "fee_amount": "$150"},
        },
        "dapproval": {
            "3227": {"approver_name": "Renee Shaw", "approval_date": "2026-09-27", "policy_version": "v4.2", "legal_notes": "Approved subject to KYC verification.", "decision_summary": "Customer meets policy thresholds.", "fee_type": "Approval Fee", "fee_amount": "$150"},
            "9001": {"approver_name": "Devon White", "approval_date": "2026-10-05", "policy_version": "v4.3", "legal_notes": "Rental product exception reviewed and approved.", "decision_summary": "Approval on expedited review track.", "fee_type": "Expedite Fee", "fee_amount": "$90"},
        },
    }

    with Session(engine) as session:
        session.execute(delete(ClientTemplateMap))
        session.commit()

        existing = select(ClientTemplateMap.client_id, ClientTemplateMap.template_name)
        rows = session.execute(existing).all()
        row_keys = {(client_id, template_name) for client_id, template_name in rows}

        for template_name, clients in template_clients.items():
            docs_dir = TEMPLATE_ROOT / template_name / "docs"
            vars_dir = TEMPLATE_ROOT / template_name / "variables"
            for client_id in clients:
                doc_filename = f"client_{client_id}.docx"
                doc_path = docs_dir / doc_filename
                _create_template_doc(doc_path, template_name)
                if (client_id, template_name) not in row_keys:
                    session.add(ClientTemplateMap(client_id=client_id, template_name=template_name, docx_filename=doc_filename))

                versions = [1, 2]
                for version in versions:
                    xml_path = vars_dir / f"client_{client_id}_{template_name}_v{version}.xml"
                    if not xml_path.exists():
                        _write_variable_xml(xml_path, template_name, client_id, version, sample_values[template_name][client_id])

        session.commit()


if __name__ == "__main__":
    seed_mock_data()
    print(f"Seeded database at {DB_PATH}")
