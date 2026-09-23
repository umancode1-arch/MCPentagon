from __future__ import annotations

from pathlib import Path

from docx import Document

from mcp_server.db import BASE_DIR, GENERATED_ROOT


def generate_document(docx_path: str, variables: dict) -> dict:
    """Populate placeholders in a template document and save a working copy under generated/."""
    source_doc = Document(docx_path)
    placeholders = []

    def replace_text(paragraph):
        if "{{" not in paragraph.text:
            return
        for key, value in variables.items():
            paragraph.text = paragraph.text.replace(f"{{{{{key}}}}}", str(value))

    for paragraph in source_doc.paragraphs:
        replace_text(paragraph)

    for table in source_doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "{{" not in cell.text:
                    continue
                for key, value in variables.items():
                    cell.text = cell.text.replace(f"{{{{{key}}}}}", str(value))

    for paragraph in source_doc.paragraphs:
        if "{{" in paragraph.text:
            placeholders.append(paragraph.text)
    for table in source_doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "{{" in cell.text:
                    placeholders.append(cell.text)

    output_path = GENERATED_ROOT / Path(docx_path).name
    source_doc.save(output_path)
    return {"path": str(output_path), "unfilled_placeholders": placeholders}
